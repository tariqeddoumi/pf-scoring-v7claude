# -*- coding: utf-8 -*-
"""Vérifie un classeur recalculé (par Excel ou LibreOffice) face au moteur de référence.

Usage :
    soffice --headless --convert-to xlsx --outdir /tmp/recalc Outil_Scoring_PF_V8_v4.xlsx
    python verifier.py /tmp/recalc/Outil_Scoring_PF_V8_v4.xlsx
    python verifier.py --lot2 /tmp/recalc/Outil_Scoring_PF_V8_v4_lot2.xlsx   (variante « Lot 2 actif »,
                                                                         python build_tool.py --sans-vba --lot2)

LibreOffice doit recalculer au chargement (Outils › Options › Calc › Formule : « Toujours
recalculer »), sinon les valeurs lues sont celles du cache, vides pour un fichier généré.

Contrôles : pour chacun des dix-huit dossiers, score, notes, plafonds C1 à C17, verrous,
régime, statut, noyau N1 à N6, indicateurs et volet climatique (section 14) identiques au
moteur ; aucune saisie obligatoire manquante ; aucune valeur d'erreur dans le classeur ;
concordance à 18 / 18.
"""
import sys
import openpyxl
import build_tool as bt
import cas_tous
import moteur_reference as mr

ERREURS = ("#VALUE!", "#NAME?", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")


def proche(a, b, tol=1e-6):
    return isinstance(a, (int, float)) and abs(a - b) <= tol


def verifier(chemin, lot2=False):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    R = bt.R
    pb = []

    # 1. aucune valeur d'erreur
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERREURS):
                    pb.append(f"{ws.title}!{c.coordinate} : {c.value}")

    # 2. chaque dossier face au moteur
    for d in cas_tous.CAS:
        ws = wb[d["onglet"]]
        r = mr.evaluer(d, lot2=lot2)
        ec = mr.evaluer_climat(d["climat"], d["secteur"], r["periodes"])
        v = lambda adr: "" if ws[adr].value is None else ws[adr].value
        k = d["onglet"]

        def egal(lib, lu, attendu, tol=1e-6):
            ok = proche(lu, attendu, tol) if isinstance(attendu, float) else lu == attendu
            if not ok:
                pb.append(f"{k} · {lib} : lu {lu!r}, attendu {attendu!r}")

        egal("score", v(f"B{R['global']}"), float(r["score"]), 1e-9)
        egal("note calculée", v(f"B{R['note_calc']}"), r["note_calculee"])
        egal("note approuvée", v(f"B{R['na']}"), r["note_approuvee"] if r["note_approuvee"] is not None else "—")
        egal("plafond", v(f"B{R['plafond']}"), r["plafond"])
        egal("régime calculé", v(f"B{R['prop_base']}"), r["regime_calcule"])
        egal("régime retenu", v(f"B{R['prop']}"), r["regime_retenu"])
        mots = lambda t: " ".join((t or "").replace(",", " ").split())
        egal("statut", mots(v(f"B{R['statut']}")), mots(r["statut"]))
        egal("verrous", v("B9"), " ".join(r["verrous_actifs"]) or "Aucun")
        egal("DSCR minimum", v(f"D{R['dmin']}"), float(r["dscr_min"]))
        egal("DSCR stressé", v(f"D{R['dstress']}"), float(r["dscr_min_s"]))
        egal("couverture", v(f"B{R['couv']}"), float(r["couverture"]))
        egal("quote-part", v(f"B{R['quote']}"), float(r["quote_part"]))
        egal("D7", v(f"D{R['d7']}"), float(r["D7"]))
        egal("complétude", v(f"B{R['compl_crit']}"), float(r["completude"]))
        for i, n in enumerate(r["noyau"]):
            lu = v(f"B{R['n1'] + i}")
            egal(f"N{i + 1}", lu, n if n is not None else "")
        for i, (c, *_x) in enumerate(bt.PLAFONDS):
            egal(f"plafond {c}", v(f"E{R['c_first'] + i}"), r["plafonds"][c])
        for i, (dom, _l, _p) in enumerate(bt.DOMAINES):
            att = r["D7"] if dom == "D7" else r["domaines"][dom]
            lu = v(f"B{R['agr_first'] + i}")
            if att is None:
                egal(f"score {dom}", lu, "")
            else:
                egal(f"score {dom}", lu, float(att), 1e-9)
        # volet climatique (section 14)
        expo = "" if ec["expo_deduite"] is None else ("Oui" if ec["expo_deduite"] else "Non")
        egal("exposition déduite", v(f"B{R['clim_expo']}"), expo)
        egal("aléas élevés sans adaptation", v(f"B{R['clim_sans']}"), len(ec["sans_adapt"]))
        egal("niveau de transition", v(f"B{R['niv_trans']}"), ec["niveau_transition"] or "")
        egal("DSCR climatique", v(f"B{R['dclim']}"), float(ec["dscr_climatique"]) if ec["dscr_climatique"] is not None else "")
        egal("exposition retenue", v(f"B{R['expo_ret']}"), r["climat_eleve_retenu"] or "")
        if lot2:
            egal("note D9.5", v(f"B{R['note_d95']}"), r["note_d95"] if r["note_d95"] is not None else "Info. insuffisante")
        ra = R["a_saisir"]
        egal("saisies manquantes", v(f"B{ra}"), 0)
        if v(f"B{ra}") != 0:
            pb.append(f"{k} · {v(f'C{ra}')}")

    # 3. concordance
    con = wb["Concordance"]
    if con["B4"].value != f"{len(cas_tous.CAS)} / {len(cas_tous.CAS)}":
        pb.append(f"Concordance : {con['B4'].value}")

    # 4. modèle vierge : calculs vides, sans erreur, saisies attendues signalées
    mod = wb["Modèle"]
    manq = mod[f"B{R['a_saisir']}"].value
    if not isinstance(manq, (int, float)) or manq <= 0:
        pb.append(f"Modèle : saisies manquantes = {manq!r}")
    return pb


if __name__ == "__main__":
    probs = verifier(sys.argv[-1], lot2="--lot2" in sys.argv)
    for p in probs:
        print("ÉCART  " + p)
    print(f"{len(probs)} écart(s) sur {len(cas_tous.CAS)} dossiers")
    sys.exit(1 if probs else 0)
